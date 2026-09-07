"""Build the Chinese research report for the v7 fixed-skill study."""

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


FONT = "STSong"
INK = RGBColor(15, 23, 42)
BLUE = RGBColor(37, 99, 235)
GRAY = RGBColor(71, 85, 105)
ORDER = ("no_skill", "oracle_v5", "oracle_v7")
LABELS = {"no_skill": "No Skill", "oracle_v5": "Oracle v5", "oracle_v7": "Oracle v7"}


def set_font(run, size=10.5, bold=None, color=None):
    run.font.name = FONT
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(key), FONT)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def border(parent, edge, value, size="8", color="000000"):
    tag = "w:tblBorders" if parent.tag.endswith("tblPr") else "w:tcBorders"
    borders = parent.first_child_found_in(tag)
    if borders is None:
        borders = OxmlElement(tag)
        parent.append(borders)
    item = borders.find(qn(f"w:{edge}"))
    if item is None:
        item = OxmlElement(f"w:{edge}")
        borders.append(item)
    item.set(qn("w:val"), value)
    item.set(qn("w:sz"), size)
    item.set(qn("w:color"), color)


def table_geometry(table, widths):
    table.autofit = False
    table.alignment = 1
    border(table._tbl.tblPr, "top", "single", "12")
    border(table._tbl.tblPr, "bottom", "single", "12")
    for edge in ("left", "right", "insideH", "insideV"):
        border(table._tbl.tblPr, edge, "nil", "0")
    for grid, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
        grid.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width / 1440)
            tc_width = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            if tc_width is None:
                tc_width = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_width)
            tc_width.set(qn("w:w"), str(width))
            tc_width.set(qn("w:type"), "dxa")


def set_cell(cell, text, header=False, size=8, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    set_font(paragraph.add_run(str(text)), size=size, bold=header, color=INK)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if header:
        border(cell._tc.get_or_add_tcPr(), "bottom", "single", "8")


def add_table(doc, headers, rows, widths, size=8):
    table = doc.add_table(rows=1, cols=len(headers))
    table_geometry(table, widths)
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    header_repeat = OxmlElement("w:tblHeader")
    header_repeat.set(qn("w:val"), "true")
    header_properties.append(header_repeat)
    for cell, text in zip(table.rows[0].cells, headers):
        set_cell(cell, text, header=True, size=size)
    for values in rows:
        row = table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for index, (cell, value) in enumerate(zip(row.cells, values)):
            set_cell(cell, value, size=size, align=WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER)
    return table


def heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(6)
    set_font(paragraph.add_run(text), size=15 if level == 1 else 12.5, bold=True, color=BLUE if level == 1 else INK)


def body(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.04
    set_font(paragraph.add_run(text), size=10.2)
    return paragraph


def bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.05
    set_font(paragraph.add_run(text), size=10)


def caption(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(6)
    set_font(paragraph.add_run(text), size=8.5, color=GRAY)


def add_figure(doc, path, caption_text, width=6.2):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    caption(doc, caption_text)


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText"); instruction.set(qn("xml:space"), "preserve"); instruction.text = " PAGE "
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, end))
    set_font(run, size=8, color=GRAY)


def percent_change(new, old):
    return 100 * (new / old - 1)


def summary_table_rows(summaries):
    rows = []
    for condition in ORDER:
        row = summaries[condition]
        rows.append((
            LABELS[condition],
            f'{100*row["task_success_rate"]:.1f}±{100*row["success_rate_repeat_std"]:.1f}%',
            f'{row["average_input_tokens"]:,.0f}', f'{row["average_output_tokens"]:,.0f}',
            f'{row["average_total_tokens"]:,.0f}', f'{row["average_latency_ms"]/1000:.2f}',
            f'${row["cost_per_successful_task_usd"]:.6f}', f'{row["average_tool_calls"]:.1f}',
        ))
    return rows


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("analysis", type=Path)
    parser.add_argument("--v5-analysis", type=Path, required=True)
    parser.add_argument("--figures-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    data = json.loads(args.analysis.read_text(encoding="utf-8"))
    v5 = json.loads(args.v5_analysis.read_text(encoding="utf-8"))
    test = {row["condition"]: row for row in data["condition_summary"]}
    dev = {row["condition"]: row for row in data["development_condition_summary"]}
    paired = {row["condition"]: row for row in data["paired_comparisons"]}
    families = {(row["condition"], row["task_family"]): row for row in data["family_summary"]}
    v5_summary = {row["condition"]: row for row in v5["condition_summary"]}

    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.62)
    section.left_margin = section.right_margin = Inches(0.82)
    section.header_distance = section.footer_distance = Inches(0.42)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        doc.styles[style_name].font.name = FONT
        doc.styles[style_name]._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(header.add_run("FixedSkillLab | ALFWorld 固定 Skill 实验"), size=8, color=GRAY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run("研究报告 v7  |  "), size=8, color=GRAY)
    add_page_number(footer)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    set_font(title.add_run("固定 Skill 能否提高 LLM Agent 的任务质量与推理效率？"), size=20, bold=True, color=INK)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(13)
    set_font(subtitle.add_run("基于 DeepSeek-v4-flash 与 ALFWorld 的分阶段实验报告"), size=11, color=GRAY)

    heading(doc, "技术摘要")
    body(doc, "第一轮正式实验没有证明固定 Skill 有效。30 道平衡任务、5 个条件、3 次重复共 450 个 episode 中，All Skills 的成功率最高，但相对 No Skill 只高 1.11 个百分点，任务级配对置信区间覆盖 0；No Skill 的 token 和成本最低。")
    body(doc, "在开发集重写 Oracle Skill 后，v7 在独立定向测试中取得 30/30，No Skill 为 29/30，旧版 Oracle v5 为 27/30。v7 同时比 No Skill 少用 11.8% 的 Total Tokens、降低 17.4% 的延迟，并将单次成功成本降低 27.4%。这些是有利的点估计，但 10 道任务不足以证明成功率差异显著。")
    body(doc, "当前最稳妥的结论是：旧版固定 Skill 没有普遍收益；经过针对性修改的 v7 在目标类别上显示了质量与成本同时改善的信号。下一步应冻结更大、类别平衡的新测试集，而不是继续修改这 10 道题。")

    heading(doc, "1. 研究问题与实验路径")
    body(doc, "研究问题有两个：固定 Skill 是否提高任务成功率；如果成功率相近，它是否减少模型输入、输出、交互步数、时间和成本。所有成功标签均由 ALFWorld 环境返回，模型不参与评分。")
    add_table(doc, ("阶段", "数据用途", "任务与条件", "允许的结论"), (
        ("v5 正式实验", "独立测试", "30 题×5 条件×3 次", "评估初版 Skill 与 Router"),
        ("v7 开发实验", "开发与改写", "20 题×3 条件×3 次", "发现问题、修改 Skill，不作最终证明"),
        ("v7 定向测试", "独立测试", "10 题×3 条件×3 次", "检验修改后的目标机制"),
    ), (1500, 1700, 2200, 3600), 8.2)
    caption(doc, "表1  三阶段证据的用途。开发集与测试集分开解释。")

    heading(doc, "2. 初版正式实验：Skill 没有形成稳定收益")
    v5_rows = []
    for condition in ("no_skill", "body_aware_retrieval", "bm25_body_retrieval", "oracle_skill", "all_skills"):
        row = v5_summary[condition]
        v5_rows.append((row["label"], f'{100*row["task_success_rate"]:.2f}%', f'{row["average_total_tokens"]:,.0f}',
                        f'{row["average_latency_ms"]/1000:.2f}', f'${row["cost_per_successful_task_usd"]:.6f}'))
    add_table(doc, ("方法", "成功率", "Total Tokens", "延迟(s)", "成本/成功"), v5_rows, (2200, 1500, 1800, 1500, 2300), 8.3)
    caption(doc, "表2  v5 正式测试总体结果。每种条件 90 个 episode。")
    body(doc, "All Skills 为 98.89%，No Skill 为 97.78%。二者只差一次成功，配对重采样的 95% 置信区间为 -3.33 至 6.67 个百分点。All Skills 的平均 Total Tokens 是 No Skill 的 3.33 倍。这个结果支持一个负结论：把 Skill 放进上下文会增加负担，但初版 Skill 没有可证实的成功率收益。")

    heading(doc, "3. v7 开发结果：更好的 Oracle 出现，但这不是最终证据")
    add_table(doc, ("条件", "成功率\n均值±SD", "Input", "Output", "Total", "延迟(s)", "成本/成功", "工具调用"),
              summary_table_rows(dev), (1300, 1300, 1250, 1000, 1250, 1000, 1450, 950), 7.3)
    caption(doc, "表3  v7 开发集结果。20 个 valid_seen 任务×3 次重复，每种条件 60 个 episode。")
    body(doc, "Oracle v7 的开发集成功率为 98.33%，高于 No Skill 的 93.33% 和 Oracle v5 的 95.00%。但 v7 是根据这批任务暴露的问题改写出来的，所以这张表只能说明设计方向值得测试，不能当作泛化证据。")

    doc.add_page_break()
    heading(doc, "4. 独立定向测试：v7 的质量与成本点估计同时改善")
    add_table(doc, ("条件", "成功率\n均值±SD", "Input", "Output", "Total", "延迟(s)", "成本/成功", "工具调用"),
              summary_table_rows(test), (1300, 1300, 1250, 1000, 1250, 1000, 1450, 950), 7.3)
    caption(doc, "表4  v7 独立定向测试。10 个 valid_unseen 任务×3 次重复，每种条件 30 个 episode。")
    add_figure(doc, args.figures_dir / "v7_method_profile.png", "图1  独立定向测试的成功率与四项推理开销。", 6.1)
    body(doc, f'相对 No Skill，Oracle v7 的成功率高 {100*(test["oracle_v7"]["task_success_rate"]-test["no_skill"]["task_success_rate"]):.2f} 个百分点；Input Tokens 变化 {percent_change(test["oracle_v7"]["average_input_tokens"], test["no_skill"]["average_input_tokens"]):.1f}%，Output Tokens 变化 {percent_change(test["oracle_v7"]["average_output_tokens"], test["no_skill"]["average_output_tokens"]):.1f}%，Total Tokens 变化 {percent_change(test["oracle_v7"]["average_total_tokens"], test["no_skill"]["average_total_tokens"]):.1f}%，延迟变化 {percent_change(test["oracle_v7"]["average_latency_ms"], test["no_skill"]["average_latency_ms"]):.1f}%，单次成功成本变化 {percent_change(test["oracle_v7"]["cost_per_successful_task_usd"], test["no_skill"]["cost_per_successful_task_usd"]):.1f}%。')
    add_figure(doc, args.figures_dir / "v7_quality_cost_frontier.png", "图2  独立定向测试的质量—成本位置。Oracle v7 位于左上方向。", 5.55)
    body(doc, "在这批任务上，Oracle v7 没有用更长上下文换成功率，而是同时减少了 token、延迟、工具调用和成本。Oracle v5 的均值被一次 100 步、749,380 token 的失败轨迹显著拉高；因此成本优势不能只看均值，还要结合逐任务结果。")

    heading(doc, "5. 不确定性：方向有利，证据量仍小")
    pair_rows = []
    for condition in ("oracle_v5", "oracle_v7"):
        row = paired[condition]
        pair_rows.append((LABELS[condition], f'{100*row["success_rate_difference"]:.2f}',
                          f'{100*row["success_rate_ci_low"]:.2f} 至 {100*row["success_rate_ci_high"]:.2f}',
                          f'{row["success_bootstrap_p_value"]:.3f}',
                          f'{row["total_tokens_difference"]:,.0f}',
                          f'{row["total_tokens_ci_low"]:,.0f} 至 {row["total_tokens_ci_high"]:,.0f}'))
    add_table(doc, ("相对 No Skill", "成功率差\n(pp)", "成功率差 95% CI", "p", "Total 差", "Total 差 95% CI"),
              pair_rows, (1450, 1050, 2100, 700, 1300, 2900), 7.8)
    caption(doc, "表5  按任务成组、保留三次重复的配对 bootstrap（20,000 次）。")
    body(doc, "Oracle v7 的成功率差为 +3.33 个百分点，区间为 0 至 10 个百分点；双侧 bootstrap p=0.692。Total Tokens 的平均差为 -4,278，但区间跨过 0。样本量和任务异质性使区间较宽，不能把点估计写成统计显著的普遍提升。")

    heading(doc, "6. 机制证据：v7 修复了目标物体绑定")
    add_figure(doc, args.figures_dir / "v7_task_success_heatmap.png", "图3  每个任务三次重复的成功次数。差异集中在任务 113 与 114。", 6.2)
    family_rows = []
    for family, family_label in (("heat_then_place", "加热后放置"), ("two_object_place", "双物体放置")):
        for condition in ORDER:
            row = families[(condition, family)]
            family_rows.append((family_label, LABELS[condition], f'{int(row["successes"])}/{int(row["episodes"])}',
                                f'{row["average_total_tokens"]:,.0f}', f'{row["average_tool_calls"]:.1f}'))
    add_table(doc, ("任务类别", "条件", "成功", "平均 Total", "工具调用"), family_rows,
              (2100, 1700, 1200, 2200, 1800), 8.1)
    caption(doc, "表6  独立定向测试的类别分解。加热任务 8 道，双物体任务 2 道；每题重复 3 次。")
    add_figure(doc, args.figures_dir / "v7_family_profile.png", "图4  两类任务的成功率与平均 Total Tokens。", 5.6)
    body(doc, "四次失败全部出现了目标物体混淆：任务要求 cup，轨迹却选择了 mug。其中 Oracle v5 有一次耗尽 100 步。Oracle v7 在任务 113 和 114 上均为 3/3，且整批测试无失败。这个轨迹证据与 v7 的设计目标一致，但它仍是窄机制证据：测试集有 8 道加热任务，不能外推到清洗、冷却、照明和一般放置任务。")

    heading(doc, "7. 指标定义与复现规则")
    for text in (
        "Task Success Rate = 成功 episode 数 / episode 总数。成功由 ALFWorld 的 won=true 判定。",
        "Input Tokens 为全部模型调用输入 token 之和；Output Tokens 为全部模型输出 token 之和；Total Tokens = Input + Output。",
        "End-to-End Latency 覆盖路由、模型调用、环境交互和完成判定，单位为毫秒，报告中换算为秒。",
        "Cost per Successful Task = 该条件全部 episode 的模型成本 / 成功 episode 数。失败消耗也计入成本。",
        "Efficiency per 1k Tokens = Success Rate / Average Total Tokens × 1000。工具调用数是补充过程指标。",
    ):
        bullet(doc, text)
    body(doc, "三个条件使用同一 DeepSeek-v4-flash、temperature=0、同一 Agent 协议和任务顺序。Oracle v5 与 Oracle v7 只改变加载的固定 Skill 版本；No Skill 不加载 Skill。独立测试文件已冻结并计算 SHA-256。")

    heading(doc, "8. 结论边界与下一步")
    body(doc, "现有证据分成两层。第一层已经确定：初版 Skill 库没有普遍优势，过量 Skill 会显著增加上下文成本。第二层仍待确认：v7 在定向独立任务上出现了 100% 成功率和更低成本，但样本太小、类别不平衡，统计区间仍宽。")
    for text in (
        "冻结 v7，不再根据这 10 道独立测试题改写 Skill。",
        "从尚未使用的 valid_unseen 任务中预先冻结一个更大且类别平衡的确认集，继续采用同题同次配对。",
        "主比较设为 Oracle v7 对 No Skill；成功率是主指标，Input、Output、Total Tokens、延迟和单次成功成本是共同报告的成本指标。",
        "报告任务级 bootstrap 置信区间、逐任务失败矩阵与中位数，避免少数超长轨迹控制均值。",
    ):
        bullet(doc, text)
    body(doc, "如果确认集仍显示质量和成本同时改善，才能把结论从“目标机制有效信号”推进到“固定 Skill 在 ALFWorld 上具有可重复收益”。如果成功率回落，v7 仍可保留为一次有解释力的失败分析案例。")

    heading(doc, "数据与版本")
    manifest = data["manifest"]
    body(doc, f'独立测试运行：{manifest["test_run_id"]}；冻结记录：{manifest["test_records"]}；开发运行：{manifest["development_run_id"]}；开发记录：{manifest["development_records"]}。')
    body(doc, f'独立测试 SHA-256：{manifest["frozen_test_sha256"]}。分析表、失败轨迹摘要与图件位于 results/formal_v7。')

    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)


if __name__ == "__main__":
    main()
