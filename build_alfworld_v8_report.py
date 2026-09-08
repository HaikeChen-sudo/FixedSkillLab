"""Build the Chinese v8 ALFWorld confirmation report with true three-line tables."""

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from build_alfworld_v7_report import (
    GRAY, INK, LABELS, ORDER, add_figure, add_page_number, add_table,
    body, bullet, caption, heading, percent_change, set_font,
)


def report_summary_rows(summary):
    rows = []
    for condition in ORDER:
        row = summary[condition]
        rows.append((
            LABELS[condition], f'{int(row["successes"])}/{int(row["episodes"])}',
            f'{100*row["task_success_rate"]:.2f}%', f'{row["average_input_tokens"]:,.0f}',
            f'{row["average_output_tokens"]:,.0f}', f'{row["average_total_tokens"]:,.0f}',
            f'{row["average_latency_ms"]/1000:.2f}', f'${row["cost_per_successful_task_usd"]:.6f}',
        ))
    return rows


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("analysis", type=Path)
    parser.add_argument("--v7-analysis", type=Path, required=True)
    parser.add_argument("--v5-analysis", type=Path, required=True)
    parser.add_argument("--figures-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    data = json.loads(args.analysis.read_text(encoding="utf-8"))
    v7 = json.loads(args.v7_analysis.read_text(encoding="utf-8"))
    v5 = json.loads(args.v5_analysis.read_text(encoding="utf-8"))
    summary = {row["condition"]: row for row in data["condition_summary"]}
    paired = {row["condition"]: row for row in data["paired_comparisons"]}
    families = {(row["condition"], row["task_family"]): row for row in data["family_summary"]}
    directions = {(row["condition"], row["task_family"], row["metric"]): row for row in data["paired_directions"]}
    task_diffs = data["task_differences"]
    v7_summary = {row["condition"]: row for row in v7["condition_summary"]}
    v5_summary = {row["condition"]: row for row in v5["condition_summary"]}

    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.64)
    section.left_margin = section.right_margin = Inches(0.72)
    section.header_distance = section.footer_distance = Inches(0.38)
    normal = doc.styles["Normal"]
    normal.font.name = "STSong"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "STSong")
    normal.font.size = Pt(10.2)
    normal.paragraph_format.space_after = Pt(4)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = doc.styles[style_name]
        style.font.name = "STSong"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "STSong")
        style.font.color.rgb = INK

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(header.add_run("FixedSkillLab | ALFWorld 固定 Skill 实验"), size=8, color=GRAY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run("研究报告 v8  |  "), size=8, color=GRAY)
    add_page_number(footer)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(5)
    set_font(title.add_run("固定 Skill 在 ALFWorld 中的复现边界"), size=20, bold=True, color=INK)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    set_font(subtitle.add_run("DeepSeek-v4-flash 的分阶段实验与 v8 确认结果"), size=11, color=GRAY)

    heading(doc, "技术摘要")
    body(doc, "v8 没有复现 v7 定向测试中的效率优势。20 道与开发集不重叠的 valid_seen 任务、3 个条件、3 次重复共 180 个 episode 中，No Skill 与 Oracle v7 均为 60/60；Oracle v7 的平均 Total Tokens 高 33.3%，端到端延迟高 16.5%，单次成功成本高 15.0%。")
    body(doc, "配对 bootstrap 的成本区间均跨过 0，不能据此断言 Skill 必然增加成本。可以确认的是，v8 不支持“v7 在保持成功率的同时稳定降低推理开销”这一说法。聚合结果上，No Skill 在质量相同的情况下占用更少 token、时间和成本。")
    body(doc, "逐任务差异很大。Oracle v7 在 23/60 个配对 episode 中减少 Total Tokens，在 37/60 个中增加；差值中位数只有 +2,413，均值却是 +15,358。少数超长轨迹控制了均值，因此后续研究必须同时报告均值、中位数、任务级区间和逐题轨迹。")

    heading(doc, "1. 证据链与确认问题")
    body(doc, "本项目先在 Mini Lab 验证 Agent、工具调用和自动评分，再进入 ALFWorld。v5 检验初版固定 Skill 与检索方法；v7 根据开发集暴露的目标绑定与动作语义问题改写两个 Oracle Skill；v8 用未参与开发的任务检查这一改写能否复现。所有成功标签都来自 ALFWorld 的 won 字段，模型不参与评分。")
    add_table(doc, ("阶段", "数据角色", "规模", "主要结论"), (
        ("v5 正式实验", "独立测试", "30题×5条件×3次", "初版 Skill 无稳定收益；All Skills 成本高"),
        ("v7 定向测试", "窄范围独立测试", "10题×3条件×3次", "v7 出现质量与成本同时改善的信号"),
        ("v8 确认实验", "任务不重叠确认集", "20题×3条件×3次", "成功率持平，效率优势未复现"),
    ), (1500, 1900, 2100, 4100), 8.2)
    caption(doc, "表1  三阶段数据用途。v8 与开发任务不重叠，但仍属于 valid_seen。")

    heading(doc, "2. v8 的六项核心指标")
    add_table(doc, ("条件", "成功", "成功率", "Input", "Output", "Total", "延迟(s)", "成本/成功"),
              report_summary_rows(summary), (1300, 900, 1100, 1250, 1100, 1250, 1050, 1750), 7.5)
    caption(doc, "表2  v8 总体结果。每个条件 60 个 episode；成本包含失败 episode 的消耗。")
    add_figure(doc, args.figures_dir / "v8_six_metric_profile.png", "图1  Task Success Rate、Input Tokens、Output Tokens、Total Tokens、End-to-End Latency 与 Cost per Successful Task。", 6.6)
    body(doc, f'Oracle v7 相对 No Skill：Input Tokens {percent_change(summary["oracle_v7"]["average_input_tokens"], summary["no_skill"]["average_input_tokens"]):+.1f}%，Output Tokens {percent_change(summary["oracle_v7"]["average_output_tokens"], summary["no_skill"]["average_output_tokens"]):+.1f}%，Total Tokens {percent_change(summary["oracle_v7"]["average_total_tokens"], summary["no_skill"]["average_total_tokens"]):+.1f}%，延迟 {percent_change(summary["oracle_v7"]["average_latency_ms"], summary["no_skill"]["average_latency_ms"]):+.1f}%，单次成功成本 {percent_change(summary["oracle_v7"]["cost_per_successful_task_usd"], summary["no_skill"]["cost_per_successful_task_usd"]):+.1f}%。两者成功率相同。')
    add_figure(doc, args.figures_dir / "v8_quality_cost_frontier.png", "图2  v8 的质量与成本位置。No Skill 与 Oracle v7 同为 100%，No Skill 成本更低。", 5.45)

    doc.add_page_break()
    heading(doc, "3. 配对统计没有支持稳定优势")
    pair_rows = []
    for condition in ("oracle_v5", "oracle_v7"):
        row = paired[condition]
        pair_rows.append((
            LABELS[condition], f'{100*row["success_rate_difference"]:+.2f}',
            f'{100*row["success_rate_ci_low"]:.2f} 至 {100*row["success_rate_ci_high"]:.2f}',
            f'{row["success_bootstrap_p_value"]:.3f}', f'{row["total_tokens_difference"]:+,.0f}',
            f'{row["total_tokens_ci_low"]:,.0f} 至 {row["total_tokens_ci_high"]:,.0f}',
            f'{row["latency_ms_difference"]/1000:+.2f}',
            f'{row["cost_usd_per_episode_difference"]:+.6f}',
        ))
    add_table(doc, ("相对 No Skill", "成功率差(pp)", "成功率差 95% CI", "p", "Total 差", "Total 差 95% CI", "延迟差(s)", "单集成本差"),
              pair_rows, (1350, 1100, 1900, 650, 1150, 2350, 1000, 1350), 7.2)
    caption(doc, "表3  按任务成组、保留三次重复的配对 bootstrap，20,000 次重采样。")
    body(doc, "Oracle v7 的成功率差为 0，区间也是 0，这是两组都达到 60/60 的结果，不代表两种方法等价。Total Tokens、延迟和单个 episode 成本的 95% 区间都跨过 0。样本只有 20 道任务，轨迹长度又高度偏斜，区间较宽。")

    direction_rows = []
    direction_labels = {
        "input_tokens": "Input Tokens", "output_tokens": "Output Tokens",
        "total_tokens": "Total Tokens", "latency_ms": "Latency (ms)",
        "cost_usd": "Cost (USD)", "tool_calls": "Tool Calls",
    }
    for metric in ("input_tokens", "output_tokens", "total_tokens", "latency_ms", "cost_usd", "tool_calls"):
        row = directions[("oracle_v7", "all", metric)]
        fmt = ".6f" if metric == "cost_usd" else ",.1f"
        direction_rows.append((direction_labels[metric], format(row["mean_difference"], fmt),
                               format(row["median_difference"], fmt),
                               f'{int(row["lower"])}/{int(row["equal"])}/{int(row["higher"])}'))
    add_table(doc, ("指标", "均值差", "中位数差", "v7更低/相同/更高"), direction_rows,
              (2200, 1900, 1900, 3000), 8.0)
    caption(doc, "表4  Oracle v7 减 No Skill 的 episode 级配对差。负值表示 v7 更省。")
    body(doc, "延迟的均值差为 +4.06 秒，中位数却是 -0.27 秒；Total Tokens 的均值差为 +15,358，中位数为 +2,413。均值和中位数方向或幅度不同，说明结论受长尾轨迹影响。")

    heading(doc, "4. 成本差异集中在少数任务")
    add_figure(doc, args.figures_dir / "v8_task_token_differences.png", "图3  每道任务三次重复后的平均 Total Tokens 差。正值表示 Oracle v7 更贵。", 6.45)
    ranked = sorted(task_diffs, key=lambda row: abs(row["total_tokens_difference"]), reverse=True)[:8]
    outlier_rows = [(row["task_id"], "加热" if row["task_family"] == "heat_then_place" else "双物体",
                     f'{row["total_tokens_difference"]:+,.0f}', f'{row["latency_ms_difference"]/1000:+.2f}',
                     f'{row["cost_usd_difference"]:+.6f}') for row in ranked]
    add_table(doc, ("任务", "类别", "Total 差", "延迟差(s)", "成本差"), outlier_rows,
              (1900, 1400, 1900, 1800, 1900), 8.0)
    caption(doc, "表5  按 Total Tokens 绝对差排序的前 8 道任务。")
    body(doc, "v7 在任务 013 和 139 上大幅减少 token，但在 078、136、047、120 上明显增加。单一平均数会把这些相反方向压成一个值。固定 Skill 的作用更像任务相关的行为干预，而不是稳定的全局加速器。")

    heading(doc, "5. 任务族分解显示加热 Skill 的成本问题")
    family_rows = []
    for family, label in (("heat_then_place", "加热后放置"), ("two_object_place", "双物体放置")):
        for condition in ORDER:
            row = families[(condition, family)]
            family_rows.append((label, LABELS[condition], f'{int(row["successes"])}/{int(row["episodes"])}',
                                f'{row["average_input_tokens"]:,.0f}', f'{row["average_total_tokens"]:,.0f}',
                                f'{row["average_latency_ms"]/1000:.2f}', f'${row["cost_per_successful_task_usd"]:.6f}'))
    add_table(doc, ("任务族", "条件", "成功", "Input", "Total", "延迟(s)", "成本/成功"), family_rows,
              (1600, 1300, 1000, 1450, 1450, 1200, 1800), 7.6)
    caption(doc, "表6  任务族分解。加热类 18 episodes，双物体类 42 episodes。")
    add_figure(doc, args.figures_dir / "v8_family_profile.png", "图4  两类任务的成功率、Total Tokens 与延迟。", 6.55)
    body(doc, "加热任务中三组均为 18/18，但 Oracle v7 的平均 Total Tokens 为 62,399，No Skill 为 23,281；延迟分别为 27.99 秒和 17.81 秒。双物体任务中 v7 与 No Skill 都是 42/42，成本接近。v8 的总体成本差主要来自加热任务和少数双物体长轨迹。")

    doc.add_page_break()
    heading(doc, "6. v7 信号没有跨任务集复现")
    comparison_rows = []
    for name, label, source in (("v7", "v7 定向 valid_unseen", v7_summary), ("v8", "v8 确认 valid_seen", summary)):
        no, skill = source["no_skill"], source["oracle_v7"]
        comparison_rows.append((label, f'{int(no["episodes"])} / {int(skill["episodes"])}',
                                f'{100*no["task_success_rate"]:.1f}% / {100*skill["task_success_rate"]:.1f}%',
                                f'{no["average_total_tokens"]:,.0f} / {skill["average_total_tokens"]:,.0f}',
                                f'{no["average_latency_ms"]/1000:.2f} / {skill["average_latency_ms"]/1000:.2f}',
                                f'${no["cost_per_successful_task_usd"]:.6f} / ${skill["cost_per_successful_task_usd"]:.6f}'))
    add_table(doc, ("测试集", "episodes N/S", "成功率 N/S", "Total N/S", "延迟 N/S", "成本 N/S"), comparison_rows,
              (2150, 1300, 1850, 2200, 1700, 2700), 7.3)
    caption(doc, "表7  No Skill（N）与 Oracle v7（S）在两批独立任务上的方向变化。")
    body(doc, "v7 定向测试中，Oracle v7 为 30/30，No Skill 为 29/30，v7 的 Total Tokens 和成本更低。v8 中两者都为 60/60，成本方向反转。前一批结果不是伪造或程序错误，它是对特定任务集成立、对确认集没有复现的结果。")
    body(doc, "v5 的较宽任务覆盖也没有给出稳定 Skill 优势：All Skills 虽达到 98.89%，但相对 No Skill 只高 1.11 个百分点，成功率差区间覆盖 0；No Skill 的平均 Total Tokens 和单次成功成本最低。v5、v7、v8 合在一起支持的是异质性，而非统一增益。")

    heading(doc, "7. 指标定义与实验规则")
    for text in (
        "Task Success Rate = 成功 episode 数 / episode 总数；成功由 ALFWorld 的 won=true 判定。",
        "Input Tokens 是全部模型调用输入 token 之和；Output Tokens 是全部输出 token 之和。",
        "Total Tokens = Input Tokens + Output Tokens。",
        "End-to-End Latency 覆盖路由、模型调用、环境交互和完成判定；报告中以秒表示。",
        "Cost per Successful Task = 条件内全部 episode 的模型成本 / 成功 episode 数，失败消耗计入分子。",
        "同一道题、同一次重复在不同条件之间配对；bootstrap 以任务为抽样单位，保留任务内三次重复。",
    ):
        bullet(doc, text)
    body(doc, "三个条件使用同一 DeepSeek-v4-flash、temperature=0、Agent 协议、任务顺序和最大步数。条件间只改变是否加载固定 Skill 以及 Skill 版本。冻结文件保存完整轨迹，不依赖模型自评。")

    heading(doc, "8. 局限与当前结论")
    for text in (
        "v8 只有 20 道任务，并且只覆盖加热后放置与双物体放置；100% 成功率存在天花板效应。",
        "v8 与开发任务不重叠，但属于 valid_seen，不能当作新的 OOD 证据。",
        "三次重复提高了对模型随机性的观察，不能替代更多独立任务；统计区间按任务成组计算。",
        "结果只覆盖 DeepSeek-v4-flash 和当前文字交互协议，不能外推到其他模型、视觉环境或不同 Agent。",
    ):
        bullet(doc, text)
    body(doc, "当前不应宣称固定 Skill 在 ALFWorld 上有效。更准确的表述是：人工 Skill 能在部分目标类型和任务实例上改变行为，但收益不稳定；过长或不匹配的 Procedure 会增加上下文与交互成本。Oracle 上界尚未形成可复现优势，因此现在扩展 Router 研究没有足够依据。")

    heading(doc, "9. 下一步")
    for text in (
        "冻结 v8，不再根据这 20 道确认题修改 Skill。把 v7 的正结果与 v8 的反转同时保留。",
        "转向新的标准环境做小规模 Pilot，优先选择 ScienceWorld；先跑 No Skill 与人工 Oracle，确认 Oracle 上界后再比较检索器。",
        "继续保留六项核心指标、任务级配对区间、中位数、逐题差异和失败轨迹。只有质量或成本出现可复现收益，才扩大实验规模。",
    ):
        bullet(doc, text)

    heading(doc, "数据与版本")
    manifest = data["manifest"]
    body(doc, f'运行 ID：{manifest["run_id"]}；冻结记录：{manifest["records"]}；任务数：{manifest["tasks"]}；每题重复 {manifest["repeats"]} 次。')
    body(doc, f'冻结文件 SHA-256：{manifest["frozen_sha256"]}。统计表、失败摘要和图件位于 results/formal_v8。v5 基线 No Skill 成功率为 {100*v5_summary["no_skill"]["task_success_rate"]:.2f}%。')

    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)


if __name__ == "__main__":
    main()
