"""Build the Chinese ScienceWorld v6 report with true three-line tables."""

import argparse
import json
from collections import Counter, defaultdict
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

import build_alfworld_v7_report as report_helpers
from build_alfworld_v7_report import GRAY, INK, add_table, bullet, caption, set_font


FONT = "Songti SC"
report_helpers.FONT = FONT

ORDER = ("no_skill", "all_skills", "flat_retrieval", "body_aware_retrieval", "oracle_v3")
LABELS = {
    "no_skill": "No Skill",
    "all_skills": "All Skills",
    "flat_retrieval": "Flat Retrieval",
    "body_aware_retrieval": "Body-aware Retrieval",
    "oracle_v3": "Oracle v3",
}


def heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_with_next = True
    set_font(paragraph.add_run(text), size=15 if level == 1 else 12.2, bold=True, color=INK)
    return paragraph


def body(doc, text, italic=False):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.12
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = paragraph.add_run(text)
    set_font(run, size=10.3, color=INK)
    run.italic = italic
    return paragraph


def add_figure(doc, path, caption_text, width=6.65):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    shape = paragraph.add_run().add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set("descr", caption_text)
    caption(doc, caption_text)


def percent_change(new, old):
    return 100 * (new / old - 1)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("analysis", type=Path)
    parser.add_argument("--figures-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    data = json.loads(args.analysis.read_text(encoding="utf-8"))
    summary = {row["condition"]: row for row in data["condition_summary"]}
    repeats = {row["condition"]: row for row in data["repeat_statistics"]}
    paired = {row["condition"]: row for row in data["paired_statistics"]}
    retrieval = {row["condition"]: row for row in data["retrieval_summary"]}
    failures = [
        json.loads(line)
        for line in (args.analysis.parent / "failures.jsonl").read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]

    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.68)
    section.left_margin = section.right_margin = Inches(0.72)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.3)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12
    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = doc.styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.color.rgb = INK

    title_style_ppr = doc.styles["Title"]._element.pPr
    title_border = title_style_ppr.find(qn("w:pBdr")) if title_style_ppr is not None else None
    if title_border is not None:
        title_style_ppr.remove(title_border)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(5)
    set_font(title.add_run("ScienceWorld 固定 Skill 检索的效果与成本"), size=20, bold=True, color=INK)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(11)
    set_font(subtitle.add_run("使用 deepseek-v4-flash 接口名称的五条件配对实验"), size=11, color=GRAY)
    add_table(doc, ("项目", "内容"), (
        ("运行编号", data["manifest"]["run_id"]),
        ("实验规模", "12 道任务  5 个条件  3 次重复  共 180 个 episode"),
        ("实验日期", "2026 年 9 月"),
        ("评分规则", "ScienceWorld 环境分数达到 100 记为成功"),
    ), (1900, 7900), 8.4)
    caption(doc, "表1  报告范围与评分口径")

    heading(doc, "摘要")
    body(doc, "本研究检验人工固定 Skill 的加载方式是否影响语言模型 Agent 在 ScienceWorld 文字环境中的任务成功率与推理成本。实验选取六类任务的 12 个官方 test variation，比较 No Skill、All Skills、Flat Retrieval、Body-aware Retrieval 和 Oracle v3。每个任务在每种条件下重复三次，共得到 180 个 episode。所有成功标签由环境分数计算，模型不能为自己评分。")
    body(doc, "Body-aware Retrieval 的成功率为 88.89%，No Skill 为 69.44%，差值为 19.44 个百分点。以任务为重采样单位的 Bootstrap 95% 置信区间为 8.33 至 30.56 个百分点。Body-aware 的平均 Total Tokens 为 244,086，比 No Skill 少 67.19%；平均延迟少 45.21%，每个成功任务的估算成本少 62.10%。Flat Retrieval 的成功率差只有 5.56 个百分点，置信区间跨过 0，平均 Total Tokens 反而增加 5.78%。")
    body(doc, "结果支持一个受限判断：在这组固定任务、当前 Skill 库和以 deepseek-v4-flash 为请求名称的模型服务上，检索时读取 Procedure 能提高质量并降低开销；只读取名称和描述没有稳定收益。任务级符号检验与 episode 级 McNemar 检验给出的证据强度不同，独立任务又只有 12 道，因此不能把这组结果外推到全部 ScienceWorld、其他模型或其他 Skill 库。")
    body(doc, "关键词  大语言模型 Agent  固定 Skill  Skill 检索  ScienceWorld  推理成本", italic=True)

    heading(doc, "一 研究问题与实验边界")
    body(doc, "研究问题有两个。第一，人工写定的执行 Procedure 能否提高交互任务的成功率。第二，如果多个方法取得相近质量，哪种方法消耗更少的输入 Token、输出 Token、时间和费用。本阶段不训练 Router，不自动学习 Skill，也不使用 PPO 或 Skill graph。")
    body(doc, "ScienceWorld 是面向基础科学任务的文字交互环境，Agent 需要观察场景、移动、操作物体、测量并改变物质状态 [1]。本实验覆盖 boil、melt、freeze、use-thermometer、find-living-thing 和 find-non-living-thing 六类任务，每类按固定规则选取第 3 和第 4 个官方 test variation。12 道题在 API 运行前固定，并用环境生成的 Gold path 验证可在 400 个环境步内达到 100 分。")

    heading(doc, "二 实验设计")
    heading(doc, "实验条件", level=2)
    add_table(doc, ("条件", "加载内容", "选择方式", "实验角色"), (
        ("No Skill", "无", "不检索", "普通 Agent 基线"),
        ("All Skills", "全部 6 个 v3 Skill", "不检索", "测试完整上下文的作用"),
        ("Flat Retrieval", "Top 3 Skill", "Name 与 Description", "轻量检索基线"),
        ("Body-aware Retrieval", "Top 3 Skill", "完整五字段 Skill 正文", "正文感知检索方法"),
        ("Oracle v3", "人工 Gold Skill", "任务类型直接映射", "已知正确 Skill 的上界"),
    ), (2050, 2250, 2500, 3000), 8.0)
    caption(doc, "表2  五个实验条件  条件间只改变提供给模型的 Skill")
    body(doc, "Skill 库含 6 个手工条目。每个条目统一包含 Name、Description、Activation Condition、Procedure 和 Termination Condition。v3 版本保留六个 Skill 的名称，只修订了 boil、melt 和 freeze 三类状态变化流程。Oracle v3 依据任务类型加载一个人工 Gold Skill；它不承担真实路由成本，所以只用于估计上界。")
    body(doc, "Flat 与 Body-aware 均使用 hash-embedding-v1：程序把英文词和相邻二元词组哈希到 256 维带符号向量，再用余弦相似度返回 Top 3。Flat 只向检索器提供 Name 和 Description；Body-aware 提供完整 Skill。该方法是确定性本地基线，不是预训练语义向量模型。")

    heading(doc, "运行协议", level=2)
    add_table(doc, ("项目", "固定设置"), (
        ("模型请求名", "deepseek-v4-flash（API 兼容名称）"),
        ("采样", "temperature 0  thinking disabled"),
        ("Agent 协议", "scienceworld-v2  每轮只执行一个 act 调用并读取观察"),
        ("环境", "ScienceWorld easy simplification  test split"),
        ("最大环境步", "400"),
        ("重复", "每道任务每个条件 3 次  条件顺序按固定种子打乱"),
        ("完整性", "180 条记录  180 个唯一键  0 缺失  0 重复"),
    ), (2300, 7500), 8.2)
    caption(doc, "表3  正式实验中保持不变的设置")
    body(doc, "每个 episode 保存完整动作轨迹、任务和 variation、重复编号、条件、加载 Skill、检索排名、环境分数、模型调用次数、Token、延迟与成本。冻结文件的 SHA 256 校验值写在复现信息中。")

    heading(doc, "指标与统计", level=2)
    for text in (
        "Task Success Rate 等于成功 episode 数除以 episode 总数；环境分数达到 100 才算成功。",
        "Input Tokens 和 Output Tokens 分别累加一个 episode 内全部模型调用的输入与输出；Total Tokens 等于两者之和。",
        "End-to-End Latency 从 Agent 开始运行计时，覆盖模型调用、路由、环境动作和完成判定。",
        "Cost per Successful Task 等于一个条件内所有 episode 的估算费用之和除以成功数；失败 episode 的费用保留在分子中。",
        "Efficiency per 1k Tokens 等于成功率除以平均 Total Tokens，再乘 1000。平均分、工具调用和环境步数用于解释执行过程。",
    ):
        bullet(doc, text)
    body(doc, "主比较采用同一道题、同一次重复的配对结构。成功率差的 95% 置信区间通过 20,000 次任务聚类 Bootstrap 计算，抽样时保留同一任务的三次重复。任务级精确符号检验把每道题的平均差分为改善、变差和持平；持平任务不进入二项检验。episode 级精确 McNemar 检验列作补充。任务与运行协议在正式实验前固定，但这套推断性检验在运行完成后补充，四项比较没有作多重检验校正。")

    heading(doc, "三 总体结果")
    core_rows = []
    for condition in ORDER:
        row, repeat = summary[condition], repeats[condition]
        core_rows.append((
            LABELS[condition], f'{int(row["successes"])}/{int(row["episodes"])}',
            f'{100*row["task_success_rate"]:.2f} ± {100*repeat["task_success_rate"]["std"]:.2f}',
            f'{row["average_input_tokens"]:,.0f} ± {repeat["average_input_tokens"]["std"]:,.0f}',
            f'{row["average_output_tokens"]:,.0f} ± {repeat["average_output_tokens"]["std"]:,.0f}',
            f'{row["average_total_tokens"]:,.0f} ± {repeat["average_total_tokens"]["std"]:,.0f}',
            f'{row["average_latency_ms"]/1000:.2f} ± {repeat["average_latency_ms"]["std"]/1000:.2f}',
            f'{row["cost_per_successful_task_usd"]:.6f} ± {repeat["cost_per_successful_task_usd"]["std"]:.6f}',
        ))
    add_table(doc, ("条件", "成功", "成功率\n%", "Input Tokens", "Output Tokens", "Total Tokens", "延迟\n秒", "成本每成功\nUSD"),
              core_rows, (1450, 700, 1250, 1650, 1450, 1650, 1250, 1750), 6.8)
    caption(doc, "表4  六项核心指标  数值为总体均值 ± 重复间标准差  每个条件 36 个 episode")

    supplemental_rows = []
    for condition in ORDER:
        row, repeat = summary[condition], repeats[condition]
        supplemental_rows.append((
            LABELS[condition],
            f'{row["average_score"]:.2f} ± {repeat["average_score"]["std"]:.2f}',
            f'{row["average_tool_calls"]:.2f} ± {repeat["average_tool_calls"]["std"]:.2f}',
            f'{row["average_environment_moves"]:.2f} ± {repeat["average_environment_moves"]["std"]:.2f}',
            f'{row["efficiency_per_1k_tokens"]:.6f} ± {repeat["efficiency_per_1k_tokens"]["std"]:.6f}',
        ))
    add_table(doc, ("条件", "平均分", "工具调用", "环境步数", "每千 Token 效率"), supplemental_rows,
              (2000, 1900, 1900, 1900, 2500), 7.6)
    caption(doc, "表5  补充过程指标  环境中的 wait 可一次消耗多个环境步")
    add_figure(doc, args.figures_dir / "scienceworld_v6_six_metrics.png",
               "图1  五种条件的任务成功率与五项推理开销  所有绝对量柱形图从零开始", 6.75)

    no_skill = summary["no_skill"]
    body_aware = summary["body_aware_retrieval"]
    all_skills = summary["all_skills"]
    flat = summary["flat_retrieval"]
    oracle = summary["oracle_v3"]
    body(doc, f'Body-aware Retrieval 相对 No Skill 的成功率高 {100*(body_aware["task_success_rate"]-no_skill["task_success_rate"]):.2f} 个百分点。Input Tokens 少 {-percent_change(body_aware["average_input_tokens"], no_skill["average_input_tokens"]):.2f}%，Output Tokens 少 {-percent_change(body_aware["average_output_tokens"], no_skill["average_output_tokens"]):.2f}%，Total Tokens 少 {-percent_change(body_aware["average_total_tokens"], no_skill["average_total_tokens"]):.2f}%，延迟少 {-percent_change(body_aware["average_latency_ms"], no_skill["average_latency_ms"]):.2f}%，每个成功任务的成本少 {-percent_change(body_aware["cost_per_successful_task_usd"], no_skill["cost_per_successful_task_usd"]):.2f}%。')
    body(doc, f'All Skills 的成功率为 {100*all_skills["task_success_rate"]:.2f}%，接近 Oracle v3 的 {100*oracle["task_success_rate"]:.2f}%，但平均 Total Tokens 是 Body-aware 的 {all_skills["average_total_tokens"]/body_aware["average_total_tokens"]:.2f} 倍。Flat Retrieval 只比 No Skill 高 {100*(flat["task_success_rate"]-no_skill["task_success_rate"]):.2f} 个百分点，平均 Total Tokens 高 {percent_change(flat["average_total_tokens"], no_skill["average_total_tokens"]):.2f}%。')
    add_figure(doc, args.figures_dir / "scienceworld_v6_quality_cost.png",
               "图2  任务成功率与每个成功任务成本  Oracle 是已知 Gold Skill 的上界  Body-aware 是成本最低的可检索方法", 6.15)

    heading(doc, "四 配对推断与证据强度")
    pair_rows = []
    for condition in ORDER[1:]:
        row = paired[condition]
        pair_rows.append((
            LABELS[condition], f'{100*row["success_rate_difference"]:+.2f}',
            f'{100*row["task_cluster_bootstrap_ci_low"]:.2f} 至 {100*row["task_cluster_bootstrap_ci_high"]:.2f}',
            f'{row["condition_only_successes"]}/{row["baseline_only_successes"]}',
            f'{row["exact_mcnemar_p_value"]:.3f}',
            f'{row["tasks_positive"]}/{row["tasks_negative"]}/{row["tasks_tied"]}',
            f'{row["exact_task_sign_p_value"]:.3f}',
        ))
    add_table(doc, ("相对 No Skill", "成功率差\n百分点", "任务聚类 95% CI\n百分点", "条件独有\n基线独有", "McNemar\np", "任务\n改善变差持平", "符号检验\np"),
              pair_rows, (1700, 1050, 1900, 1300, 900, 1500, 1100), 7.1)
    caption(doc, "表6  成功率配对统计  p 值为双侧且未作多重比较校正")
    body(doc, "Body-aware 在 12 道任务中改善 6 道、没有变差、6 道持平，任务级符号检验 p=0.031；任务聚类 Bootstrap 区间完全高于 0。episode 级 McNemar p=0.065，略高于常用的 0.05 阈值。三次重复属于同一任务，不能把 36 个 episode 当成 36 道独立题。报告同时保留两种检验，不用其中一个替代另一个。")
    body(doc, "All Skills 的 episode 级 McNemar p=0.021，但任务级符号检验 p=0.219。它改善 5 道、损害 1 道、其余持平。Flat Retrieval 的 Bootstrap 区间包含 0，两种精确检验也没有支持稳定差异。Oracle v3 的上界结果方向最强，但它依赖预先知道 Gold Skill，不能直接当作部署方法。")

    doc.add_page_break()
    heading(doc, "五 检索与任务层面结果")
    retrieval_rows = []
    for condition in ("flat_retrieval", "body_aware_retrieval"):
        row = retrieval[condition]
        missed = "boil v023 v024  melt v023 v024" if condition == "flat_retrieval" else "boil v023 v024"
        retrieval_rows.append((
            LABELS[condition], f'{int(row["gold_hit_at_3"])}/{int(row["episodes"])}',
            f'{100*row["gold_hit_at_3_rate"]:.2f}%', f'{100*row["success_rate_when_hit"]:.2f}%',
            f'{100*row["success_rate_when_missed"]:.2f}%', missed,
        ))
    add_table(doc, ("检索器", "Gold 命中", "Hit at 3", "命中时成功率", "漏检时成功率", "漏检任务"),
              retrieval_rows, (2100, 1200, 1100, 1500, 1500, 2600), 7.7)
    caption(doc, "表7  在线检索结果  每个 variation 重复三次  因此同一漏检任务出现三条 episode")
    body(doc, "Flat 的 Gold Hit@3 为 66.67%，Body-aware 为 83.33%。Flat 在 boil 和 melt 各漏掉两个 variation；Body-aware 只漏掉两个 boil variation。两种方法在命中 Gold Skill 时都更容易成功，Body-aware 在命中与漏检两种情况下都取得更高成功率。检索命中与成功相关，但不能单独证明因果，因为 Top 3 中其他 Skill、任务难度和后续执行都会改变结果。")
    add_figure(doc, args.figures_dir / "scienceworld_v6_task_heatmap.png",
               "图3  12 道任务在五个条件下的成功次数  每个单元格分母为 3", 6.75)
    body(doc, "差异集中在状态变化任务。No Skill 在 boil v023 三次全败，在 boil v024、melt v024 和 freeze v023 也不稳定。Body-aware 在这四道难题上都取得 2/3，其余八道均为 3/3。All Skills 唯一系统性失败是 melt v024 的 0/3，说明完整 Skill 上下文会帮助多数任务，也可能在个别长流程中形成干扰或拖长轨迹。")

    heading(doc, "六 失败分析")
    by_condition = Counter(row["condition"] for row in failures)
    by_category = Counter((row["condition"], row["category"]) for row in failures)
    failed_tasks = defaultdict(set)
    for row in failures:
        failed_tasks[row["condition"]].add(row["task_id"].replace(":v", " v"))
    failure_rows = []
    for condition in ORDER:
        failure_rows.append((
            LABELS[condition], by_condition[condition],
            by_category[(condition, "environment_move_budget_exhausted")],
            by_category[(condition, "premature_model_stop")],
            by_category[(condition, "execution_error_other")],
            "  ".join(sorted(failed_tasks[condition])) or "无",
        ))
    add_table(doc, ("条件", "失败", "动作预算耗尽", "模型提前停止", "其他执行错误", "涉及任务"), failure_rows,
              (1700, 800, 1450, 1450, 1300, 3200), 7.4)
    caption(doc, "表8  29 个失败 episode 的规则化分类")
    add_figure(doc, args.figures_dir / "scienceworld_v6_retrieval_failures.png",
               "图4  左图比较检索命中与成功率  右图按条件分解失败类型", 6.75)
    body(doc, "29 次失败中，18 次耗尽环境动作预算，占 62.07%；8 次在环境未完成时模型停止，占 27.59%；3 次归入其他执行错误。No Skill 有 11 次失败，Flat 有 9 次，Body-aware 有 4 次，Oracle v3 有 2 次。失败减少的主要来源不是评分规则变化，而是困难状态变化任务中的轨迹缩短。")
    body(doc, "动作预算耗尽也暴露了一个系统因素。Agent 把完整历史继续送入模型，长 episode 的 Input Tokens 会快速累积；ScienceWorld 的 wait 还会一次推进十个环境步。因此，平均 Input Tokens 同时反映 Skill 文本长度、模型调用次数和轨迹是否及时结束，不能把它只解释成一次性 Prompt 长度。")

    heading(doc, "七 结果解释")
    heading(doc, "正文感知检索是当前最好的可部署条件", level=2)
    body(doc, "Body-aware 在可检索方法中同时取得更高成功率、更低 Total Tokens 和更低成功成本。它读取 Procedure 后将 Gold Hit@3 从 Flat 的 66.67% 提高到 83.33%，并把平均工具调用从 No Skill 的 69.64 次降到 38.69 次。现有证据更符合两段机制：Procedure 改善 Skill 排名，加载后的步骤约束又减少无效探索。实验没有把路由和执行分开随机化，所以这仍是机制解释，不是独立因果估计。")
    heading(doc, "全部加载的质量较高但上下文代价仍大", level=2)
    body(doc, "All Skills 取得 91.67% 成功率，却平均使用 657,441 Total Tokens，是 Body-aware 的 2.69 倍。它相对 No Skill 的总 Token 仍然更少，是因为多数 episode 更快结束，而不是因为每次调用的 Prompt 更短。melt v024 的三次失败说明完整库也会在特定流程中产生干扰。")
    heading(doc, "Oracle 给出上界但不能替代 Router", level=2)
    body(doc, "Oracle v3 达到 94.44% 成功率和 97.11 平均分，每个成功任务的成本为 0.002701 美元。这个条件使用人工 Gold 标签，省去了真实检索错误。它说明 Skill 内容仍有改进空间，也给 Router 设定了可接近的目标；它不能用于声称系统已经实现同等的在线选择能力。")

    heading(doc, "八 局限与结论")
    for text in (
        "正式集只有 12 道独立任务。三次重复反映运行波动，不能替代更多任务；任务级检验的离散度仍高。",
        "六类任务各有两个 variation，但 boil、melt 和 freeze 的长轨迹控制了大部分 Token 与失败差异。",
        "Skill v3 在此前 variation 上经过修订。v6 任务没有参与修订，但结论仍限于同一环境和相邻任务族。",
        "API 请求使用 deepseek-v4-flash 名称、temperature 0 和 scienceworld-v2 Agent 协议。官方文档在 2026 年 9 月 14 日将该名称列为兼容名称，并说明请求由 DeepSeek-V4.1-Flash 提供服务；客户端没有冻结服务端权重。temperature 为 0 也不能排除服务侧非确定性 [2]。",
        "hash-embedding-v1 是局部哈希检索基线。结果不能代表训练过的语义 embedding、重排器或其他 Router。",
        "成本依据运行时记录的缓存命中、缓存未命中和输出 Token 估算。实验代码冻结的单价分别为每百万 Token 0.0028、0.14 和 0.28 美元；这些数值用于复现实验内比较，不等同于当前账单价格 [2]。",
        "四个条件都与 No Skill 比较，p 值未作多重比较校正。Bootstrap 与精确检验是在运行后补充，适合描述证据强度，不应包装成预注册显著性检验。",
    ):
        bullet(doc, text)
    body(doc, "这次正式实验给出了项目目前最清楚的正结果。Body-aware Retrieval 在预先固定的 ScienceWorld test variation 上比 No Skill 多成功 7 个 episode，并大幅减少平均 Token、延迟与成功成本。Flat Retrieval 没有稳定收益，说明只检索 Skill 名称和描述不足以支持复杂执行。现阶段可以写成“正文感知 Skill 检索在该固定套件上显示质量与成本的共同改善”，不能写成“Skill 对 LLM Agent 普遍有效”。")

    heading(doc, "九 复现信息")
    manifest = data["manifest"]
    add_table(doc, ("字段", "值"), (
        ("运行编号", manifest["run_id"]),
        ("套件", manifest["suite"]),
        ("冻结记录", str(manifest["records"])),
        ("冻结文件", manifest["frozen_file"]),
        ("SHA 256", manifest["sha256"]),
        ("模型", ", ".join(manifest["models"])),
        ("Agent 协议", ", ".join(manifest["agent_protocol_versions"])),
        ("成功规则", manifest["success_rule"]),
        ("分析输入", "condition_summary.csv  paired_statistics.csv  retrieval_summary.csv  failures.jsonl  task_summary.csv"),
    ), (2200, 7600), 8.0)
    caption(doc, "表9  冻结数据与分析来源")
    body(doc, "正式运行命令")
    command = (
        ".venv/scienceworld/bin/python scienceworld_experiment.py\n"
        "  --model deepseek-v4-flash --temperature 0\n"
        "  --suite scienceworld-v6-router-formal-12\n"
        "  --conditions no_skill,all_skills,flat_retrieval,body_aware_retrieval,oracle_v3\n"
        "  --max-steps 400 --repeats 3\n"
        "  --run-id scienceworld-v6-router-formal-12x3 --resume"
    )
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(command)
    set_font(run, size=8.2, color=INK)
    analysis_label = body(doc, "分析命令")
    analysis_label.paragraph_format.keep_with_next = True
    analysis_command = (
        "python3 analyze_scienceworld_formal.py results/scienceworld_results.jsonl\n"
        "  --run-id scienceworld-v6-router-formal-12x3\n"
        "  --suite scienceworld-v6-router-formal-12\n"
        "  --conditions no_skill,all_skills,flat_retrieval,body_aware_retrieval,oracle_v3\n"
        "  --repeats 3 --output-dir results/formal_scienceworld_v6"
    )
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    run = paragraph.add_run(analysis_command)
    set_font(run, size=8.2, color=INK)

    heading(doc, "参考资料")
    body(doc, "[1] Wang R, Jansen P, Côté M A, Ammanabrolu P. ScienceWorld Is your Agent Smarter than a 5th Grader. Proceedings of EMNLP 2022: 11279-11298. DOI 10.18653/v1/2022.emnlp-main.775. https://aclanthology.org/2022.emnlp-main.775/")
    body(doc, "[2] DeepSeek API Documentation. Models and Pricing. 访问日期 2026 年 9 月 14 日. https://api-docs.deepseek.com/quick_start/pricing")

    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
