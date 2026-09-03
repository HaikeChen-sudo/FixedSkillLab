"""Build the Chinese technical report for a frozen ALFWorld formal run."""

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


FONT = "Arial Unicode MS"
INK = RGBColor(11, 37, 69)
BLUE = RGBColor(46, 116, 181)
GRAY = RGBColor(89, 89, 89)
CONDITION_ORDER = (
    "no_skill", "body_aware_retrieval", "bm25_body_retrieval", "oracle_skill", "all_skills",
)
LABELS = {
    "no_skill": "No Skill",
    "body_aware_retrieval": "Body-aware",
    "bm25_body_retrieval": "BM25 Body",
    "oracle_skill": "Oracle Skill",
    "all_skills": "All Skills",
}


def set_font(run, size=11, bold=None, color=None):
    run.font.name = FONT
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn(key), FONT)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_border(parent, edge, value, size="8", color="000000"):
    tag = "w:tblBorders" if parent.tag.endswith("tblPr") else "w:tcBorders"
    borders = parent.first_child_found_in(tag)
    if borders is None:
        borders = OxmlElement(tag)
        parent.append(borders)
    border = borders.find(qn(f"w:{edge}"))
    if border is None:
        border = OxmlElement(f"w:{edge}")
        borders.append(border)
    border.set(qn("w:val"), value)
    border.set(qn("w:sz"), size)
    border.set(qn("w:color"), color)


def set_table_geometry(table, widths):
    table.autofit = False
    width = table._tbl.tblPr.first_child_found_in("w:tblW")
    if width is None:
        width = OxmlElement("w:tblW")
    if width.getparent() is None:
        table._tbl.tblPr.append(width)
    width.set(qn("w:w"), str(sum(widths)))
    width.set(qn("w:type"), "dxa")
    indent = table._tbl.tblPr.first_child_found_in("w:tblInd")
    if indent is None:
        indent = OxmlElement("w:tblInd")
    if indent.getparent() is None:
        table._tbl.tblPr.append(indent)
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")
    set_border(table._tbl.tblPr, "top", "single", "12")
    set_border(table._tbl.tblPr, "bottom", "single", "12")
    for edge in ("left", "right", "insideH", "insideV"):
        set_border(table._tbl.tblPr, edge, "nil", "0")
    for grid_col, cell_width in zip(table._tbl.tblGrid.gridCol_lst, widths):
        grid_col.set(qn("w:w"), str(cell_width))
    for row in table.rows:
        for cell, cell_width in zip(row.cells, widths):
            cell.width = Inches(cell_width / 1440)
            tc_width = cell._tc.tcPr.first_child_found_in("w:tcW")
            if tc_width is None:
                tc_width = OxmlElement("w:tcW")
            if tc_width.getparent() is None:
                cell._tc.tcPr.append(tc_width)
            tc_width.set(qn("w:w"), str(cell_width))
            tc_width.set(qn("w:type"), "dxa")


def set_cell(cell, text, header=False, size=8, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(str(text))
    set_font(run, size=size, bold=header, color=INK if header else None)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if header:
        set_border(cell._tc.get_or_add_tcPr(), "bottom", "single", "8")


def paragraph(doc, text="", bold_prefix=None):
    item = doc.add_paragraph()
    item.paragraph_format.space_after = Pt(6)
    item.paragraph_format.line_spacing = 1.10
    if bold_prefix and text.startswith(bold_prefix):
        set_font(item.add_run(bold_prefix), bold=True)
        set_font(item.add_run(text[len(bold_prefix):]))
    else:
        set_font(item.add_run(text))
    return item


def heading(doc, text, level=1):
    item = doc.add_paragraph(style=f"Heading {level}")
    item.paragraph_format.space_before = Pt(13 if level == 1 else 10)
    item.paragraph_format.space_after = Pt(8 if level == 1 else 6)
    set_font(item.add_run(text), size=16 if level == 1 else 13, bold=True, color=BLUE if level == 1 else INK)
    return item


def caption(doc, text):
    item = doc.add_paragraph()
    item.alignment = WD_ALIGN_PARAGRAPH.CENTER
    item.paragraph_format.space_before = Pt(3)
    item.paragraph_format.space_after = Pt(8)
    set_font(item.add_run(text), size=9, color=GRAY)


def bullet(doc, text):
    item = doc.add_paragraph(style="List Bullet")
    item.paragraph_format.space_after = Pt(4)
    item.paragraph_format.line_spacing = 1.10
    set_font(item.add_run(text), size=10.5)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, end))
    set_font(run, size=8.5, color=GRAY)


def main():
    parser = argparse.ArgumentParser(description="Build an ALFWorld formal report")
    parser.add_argument("analysis", type=Path)
    parser.add_argument("--figures-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    data = json.loads(args.analysis.read_text(encoding="utf-8"))
    summaries = {row["condition"]: row for row in data["condition_summary"]}
    paired = {row["condition"]: row for row in data["paired_comparisons"]}
    taxonomy = {row["category"]: row for row in data["failure_taxonomy"]}
    retrieval = data["retrieval_audit"]
    run_id = data["manifest"]["run_id"]

    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.05
    for style_name in ("Heading 1", "Heading 2"):
        doc.styles[style_name].font.name = FONT
        doc.styles[style_name]._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(header.add_run("ALFWorld 固定 Skill 路由实验 | 正式分析"), size=8.5, color=GRAY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run(f"{run_id}  |  "), size=8.5, color=GRAY)
    add_page_field(footer)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    set_font(title.add_run("ALFWorld 中固定 Skill 路由的效果与推理成本"), size=21, bold=True, color=INK)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    set_font(subtitle.add_run("DeepSeek-v4-flash · 独立测试集30题 · 5个条件 · 3次重复 · 450 episodes"), size=10.5, color=GRAY)

    heading(doc, "技术摘要")
    paragraph(doc, "目标句路由修正后，固定 Skill 仍未显示稳定收益。All Skills 的原始成功率最高（98.89%），相对 No Skill 只高1.11个百分点；任务级配对重采样的95%置信区间为-3.33至6.67个百分点。No Skill 与 BM25 Body 均为97.78%，No Skill 的推理开销最低。")
    paragraph(doc, "离线与在线 BM25 已统一使用目标句，30道任务的 Top-3 完全一致。主 Skill Hit@3 为86.7%，但检索命中没有转化为更短轨迹。主要问题更可能在 Skill 内容与执行，而不是查询构造。")
    paragraph(doc, "关键词：LLM Agent；ALFWorld；Skill Retrieval；推理成本；负结果；失败分析")

    heading(doc, "1. All Skills 原始成功率最高，No Skill 的效率最高")
    paragraph(doc, "表1列出六项核心指标。成功率同时报告三次重复的样本标准差和按任务成组重采样的95%置信区间，避免把同题的90个 episode 当成独立样本。")
    table = doc.add_table(rows=1, cols=8)
    widths = [1250, 950, 1200, 1000, 850, 1000, 900, 2210]
    set_table_geometry(table, widths)
    headers = ("方法", "成功率\n均值±SD", "95% CI", "Input\nTokens", "Output\nTokens", "Total\nTokens", "延迟\n(s)", "成本/成功\n(USD)")
    for cell, text in zip(table.rows[0].cells, headers):
        set_cell(cell, text, header=True, size=7.3)
    for condition in CONDITION_ORDER:
        row = summaries[condition]
        values = (
            LABELS[condition],
            f'{100 * row["task_success_rate"]:.2f}±{100 * row["success_rate_repeat_std"]:.2f}%',
            f'{100 * row["success_ci_low"]:.2f}–{100 * row["success_ci_high"]:.2f}%',
            f'{row["average_input_tokens"]:,.0f}',
            f'{row["average_output_tokens"]:,.0f}',
            f'{row["average_total_tokens"]:,.0f}',
            f'{row["average_latency_ms"] / 1000:.2f}',
            f'${row["cost_per_successful_task_usd"]:.6f}',
        )
        for index, (cell, value) in enumerate(zip(table.add_row().cells, values)):
            set_cell(cell, value, size=7.1, align=WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER)
    caption(doc, "表1  独立测试集上的总体结果。每种条件90个 episode；CI按30个任务成组重采样20,000次。")
    paragraph(doc, "相对 No Skill，BM25 Body 的成功率相同，但总 token 增加39.5%，单次成功成本增加16.4%。All Skills 高1.11个百分点，却增加233.0%的总 token、57.1%的单次成功成本和39.4%的延迟。该点估计不能证明这一次额外成功来自 Skill。")

    profile = args.figures_dir / "alfworld_method_profile.png"
    doc.add_picture(str(profile), width=Inches(6.2))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption(doc, "图1  五种条件的成功率、总 token、端到端延迟和单次成功成本。误差线为三次重复的样本标准差。")
    paragraph(doc, "五组成功率都高于94%，成本差异却达到数倍。All Skills 的长上下文没有形成可证实的成功率优势；BM25 Body 也没有把更准确的选择转化为更短的轨迹。")

    frontier = args.figures_dir / "alfworld_quality_cost_frontier.png"
    doc.add_picture(str(frontier), width=Inches(6.1))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption(doc, "图2  质量—成本关系。左上方向表示更高成功率和更低单次成功成本。")
    paragraph(doc, "观测到的 Pareto 前沿包含 No Skill 与 All Skills：前者成本最低，后者原始成功率高1.11个百分点。由于配对置信区间覆盖0，All Skills 的质量增益尚不能和随机波动区分。BM25 Body、Body-aware 与 Oracle 均被 No Skill 支配。")

    heading(doc, "2. 配对比较未发现 Skill 带来的成功率增益")
    paragraph(doc, "每个条件使用相同任务和重复编号，因此比较采用任务级配对重采样。差值定义为 Skill 条件成功率减去 No Skill 成功率；负值表示 Skill 条件更差。")
    paired_table = doc.add_table(rows=1, cols=4)
    set_table_geometry(paired_table, [2600, 1800, 2600, 2360])
    for cell, text in zip(paired_table.rows[0].cells, ("对比条件", "差值（百分点）", "95% CI（百分点）", "Bootstrap p")):
        set_cell(cell, text, header=True, size=8)
    for condition in ("body_aware_retrieval", "bm25_body_retrieval", "oracle_skill", "all_skills"):
        row = paired[condition]
        values = (LABELS[condition], f'{100 * row["difference"]:.2f}', f'{100 * row["ci_low"]:.2f}–{100 * row["ci_high"]:.2f}', f'{row["bootstrap_p_value"]:.3f}')
        for index, (cell, value) in enumerate(zip(paired_table.add_row().cells, values)):
            set_cell(cell, value, size=8, align=WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER)
    caption(doc, "表2  相对 No Skill 的任务级配对重采样结果。p值用于描述证据强度，不作多重比较校正。")
    paragraph(doc, "四个区间都覆盖0。All Skills 的点估计为+1.11个百分点，BM25 Body 为0，Body-aware 与 Oracle 分别为-3.33和-2.22个百分点。现有样本不能支持任何 Skill 条件优于 No Skill。")

    heading(doc, "3. 失败集中在双物体放置与加热任务")
    paragraph(doc, "Clean、Cool、Light 和 Simple Place 在所有条件下均达到100%。Two objects 的区分度最高：All Skills 为100%，No Skill、BM25 Body 与 Oracle 为86.7%，Body-aware 为73.3%。Heat 类别中 All Skills 与 Body-aware 为93.3%，Oracle 为86.7%，No Skill 与 BM25 Body 为100%。")
    family = args.figures_dir / "alfworld_family_success_heatmap.png"
    doc.add_picture(str(family), width=Inches(6.25))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption(doc, "图3  各条件在六类任务上的成功率。每个单元包含5题×3次重复，共15个 episode。")
    paragraph(doc, "这种集中性说明总体成功率已经接近天花板。当前任务集真正区分方法的是多物体状态跟踪，而不是大多数单物体任务。后续测试集应保留类别平衡，并单独报告 Two objects。")

    failure = args.figures_dir / "alfworld_failure_heatmap.png"
    doc.add_picture(str(failure), width=Inches(6.1))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption(doc, "图4  三次重复中出现失败的任务及条件。数字为失败 episode 数。")
    paragraph(doc, "14次失败分布在任务036、051、053和074。按同题同次对照，10次属于 Skill 条件失败而 No Skill 成功，2次满足“检索漏掉主 Skill且 Oracle 成功”的严格路由错误定义，另有2次为 No Skill 的其他执行错误。该分类来自配对轨迹规则，不是人工因果标注。")

    taxonomy_table = doc.add_table(rows=1, cols=3)
    set_table_geometry(taxonomy_table, [3300, 1200, 4860])
    for cell, text in zip(taxonomy_table.rows[0].cells, ("失败类别", "次数", "占全部失败")):
        set_cell(cell, text, header=True, size=8.5)
    taxonomy_labels = {
        "routing_error": "严格路由错误",
        "execution_error_other": "其他执行错误",
        "skill_condition_interference": "Skill 条件干扰或执行偏移",
    }
    for category in ("routing_error", "execution_error_other", "skill_condition_interference"):
        row = taxonomy.get(category, {"failures": 0, "share": 0})
        values = (taxonomy_labels[category], row["failures"], f'{100 * row["share"]:.1f}%')
        for index, (cell, value) in enumerate(zip(taxonomy_table.add_row().cells, values)):
            set_cell(cell, value, size=8.5, align=WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER)
    caption(doc, "表3  基于配对轨迹的互斥失败分类。")

    heading(doc, "4. 查询一致性修复没有转化为任务收益")
    paragraph(doc, f'BM25 Body 在实际运行中命中主 Gold Skill 的任务为{retrieval["primary_gold_hit_at_3_tasks"]}/30，Hit@3为{100 * retrieval["primary_gold_hit_at_3_task_rate"]:.1f}%。命中时成功率为{100 * retrieval["success_when_hit"]:.1f}%，未命中时为{100 * retrieval["success_when_missed"]:.1f}%。二者有明显相关，但不能据此断言检索命中导致成功，因为任务难度同时影响命中和执行。')
    paragraph(doc, "本轮离线与在线检索逐题核验后，Top-3 不一致数为0。BM25 Body 与 No Skill 成功率相同，却多使用39.5%的总 token。查询构造已经不是主要解释，后续分析应转向 Skill 是否提供了可执行、互不冲突且能减少搜索的程序性指导。")

    heading(doc, "5. 实验范围与指标定义")
    paragraph(doc, "实验使用 ALFWorld valid_unseen 中预先冻结的30道任务，六类任务各5道。五种条件使用同一 DeepSeek-v4-flash 模型、温度0、同一工具协议和50步上限，每题重复3次。成功仅由环境 won=true 判定，模型不参与评分。")
    for text in (
        "Task Success Rate：成功 episode 数除以该条件总 episode 数。",
        "Input Tokens：全部模型调用的输入 token；Output Tokens：模型生成 token；Total Tokens 为二者之和。",
        "End-to-End Latency：从环境重置、路由、模型调用、工具交互到结束的总时间。",
        "Cost per Successful Task：该条件全部 episode 的推理成本之和除以成功 episode 数。",
    ):
        bullet(doc, text)
    paragraph(doc, "成本按运行器中 DeepSeek-v4-flash 的缓存命中、缓存未命中和输出 token 单价计算。价格变化会改变绝对成本，但不会改变本次记录中的 token 数。")

    heading(doc, "6. 限制与稳健性")
    for text in (
        "测试集只有30道题。95%以上的总体成功率使方法差异只能由少数难题决定。",
        "三次重复仍不足以估计很小的成功率差异；配对区间包含0，结果应写成负结果而非方法排名定论。",
        "失败分类依赖同题同次对照和轨迹字段，只能识别与路由、Skill 条件或共同执行相关的模式，不能证明因果。",
        "当前 v5 test 集已经用于正式分析，后续任何 Router 或 Skill 修改都不能继续用它调参。",
    ):
        bullet(doc, text)

    heading(doc, "7. 下一步：停止在 v5 测试集上调参")
    paragraph(doc, "这套测试集已经完成正式评估，不能继续用于修改 Skill。当前结果应按负结果报告：目标句路由保证了离线与在线一致，但固定 Skill 没有产生可证实的成功率收益，并增加了推理成本。")
    paragraph(doc, "若继续研究，应在开发集上重写双物体放置和加热 Skill，明确禁止搜索与目标无关的箱子、限制重复无效动作，并把搜索预算写入终止条件。完成修改后必须冻结新的未见任务集，或更换模型和环境进行外部复验。")

    heading(doc, "数据与复现信息")
    paragraph(doc, f'冻结运行：{run_id}；记录数：{data["manifest"]["records"]}；SHA-256：{data["manifest"]["sha256"]}。原始轨迹、统计表、失败摘要和绘图数据均保存在项目 results/formal_v5 目录。')

    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
